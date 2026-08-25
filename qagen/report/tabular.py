"""Excel and CSV renderers.

One row per **step**, with the group-level columns (Test ID, Category,
Preconditions, Overall Expected Result) repeated only on the first row of each
group. That is the shape TestRail, Zephyr, and Xray import cleanly -- a row per
test case with steps crammed into one cell does not.
"""

from __future__ import annotations

import csv
from pathlib import Path

from ..models import TestSuite

HEADERS = [
    "Test ID",
    "Category",
    "Preconditions",
    "Step #",
    "Action",
    "Step Expected Result",
    "Overall Expected Result",
    "Source URL",
    "Needs Review",
]


def rows(suite: TestSuite) -> list[list[str]]:
    out: list[list[str]] = []
    for case in suite.cases:
        preconditions = " · ".join(case.preconditions)
        for index, step in enumerate(case.steps):
            first = index == 0
            out.append(
                [
                    case.test_id if first else "",
                    case.category.value if first else "",
                    preconditions if first else "",
                    str(step.step_number),
                    step.action,
                    step.expected_result,
                    case.overall_expected_result if first else "",
                    case.source_url if first else "",
                    ("yes" if case.needs_review else "") if first else "",
                ]
            )
    return out


def write_csv(suite: TestSuite, out_dir: Path) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / "test-cases.csv"
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(HEADERS)
        writer.writerows(rows(suite))
    return path


def write_xlsx(suite: TestSuite, out_dir: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / "test-cases.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "Test Cases"

    ws.append(HEADERS)
    header_fill = PatternFill("solid", start_color="1F3B57")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    review_fill = PatternFill("solid", start_color="FFF2CC")
    for row in rows(suite):
        ws.append(row)
        if row[8] == "yes":
            for cell in ws[ws.max_row]:
                cell.fill = review_fill

    widths = [12, 14, 40, 8, 52, 52, 46, 40, 13]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{ws.max_row}"

    wb.save(path)
    return path
